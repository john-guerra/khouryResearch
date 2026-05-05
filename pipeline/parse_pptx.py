"""Parse the Khoury Research Day PPTX into structured researcher records.

Reads the .pptx, extracts per-researcher: name, location, give bullets, get
bullets, primary/secondary research areas, keywords, email, links, and the
headshot photo. Writes researchers.json + photos to disk.
"""

from __future__ import annotations

import json
import re
import sys
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
PPTX_PATH = ROOT / "Khoury Research Day Give_Get_Slides.pptx"
OUT_JSON = ROOT / "pipeline" / "researchers.json"
PHOTO_DIR = ROOT / "public" / "data" / "photos"
# Files in here override auto-extracted headshots when the heuristic picks a
# non-portrait (e.g., the slide's only image is a group shot). Drop a file
# named <slug>.<ext> here and it wins.
PHOTO_OVERRIDES = ROOT / "pipeline" / "photo_overrides"

EMAIL_RE = re.compile(r"[\w\.\-+]+@[\w\.\-]+\.\w+")
URL_RE = re.compile(r"https?://\S+")
NAME_LOC_RE = re.compile(r"^(?P<name>[^()\n]+?)\s*\((?P<loc>[^)]+)\)\s*$")

GIVE_HEADERS = ("WHAT YOU CAN GIVE", "GIVE:", "GIVE")
GET_HEADERS = ("WHAT YOU WANT TO GET", "WHAT YOU CAN GET", "GET:", "GET")


@dataclass
class Researcher:
    id: str
    name: str
    location: str
    give: list[str]
    get: list[str]
    primary_area: str
    secondary_area: str
    keywords: list[str]
    email: str
    links: list[str] = field(default_factory=list)
    photo: str | None = None
    slide_index: int = 0


def slugify(name: str) -> str:
    s = re.sub(r"[^\w\s-]", "", name.lower())
    s = re.sub(r"[\s_]+", "-", s).strip("-")
    return s


def normalize_location(loc: str) -> str:
    """Strip parenthetical asides, trailing state codes, and clutter.

    "Oakland - soon!" → "Oakland", "Arlington VA" → "Arlington",
    "Portland, ME" → "Portland".
    """
    s = re.split(r"[\-,]", loc, maxsplit=1)[0].strip()
    s = re.sub(r"\s+[A-Z]{2}$", "", s)  # trailing state abbreviation
    return s.strip(" ,.")


def normalize(text: str) -> str:
    # Replace vertical tab (U+000B) with newline; collapse trailing whitespace.
    return text.replace("\x0b", "\n").strip()


def detect_section(line: str) -> str | None:
    upper = line.strip().upper().rstrip(":")
    for h in GIVE_HEADERS:
        if upper == h.rstrip(":"):
            return "give"
    for h in GET_HEADERS:
        if upper == h.rstrip(":"):
            return "get"
    return None


def parse_give_get_block(text: str) -> tuple[str | None, list[str]]:
    """A frame whose first non-empty line is GIVE/GET header; rest are bullets."""
    lines = [ln.strip() for ln in normalize(text).split("\n") if ln.strip()]
    if not lines:
        return None, []
    section = detect_section(lines[0])
    if section is None:
        return None, []
    return section, lines[1:]


KEYWORDS_LABEL = r"(?:Research\s+)?[Kk]ey\s*[Ww]ords?:"


def parse_areas_block(text: str) -> tuple[str, str, list[str]]:
    """Pull primary/secondary/keywords out of the areas frame.

    Handles two shapes:
      A) Explicit-labeled: 'Primary Research Area: X | Secondary Research Area: Y | Key words: Z'
      B) Single-line uncommitted: 'Broadening Participation / CSEd / Digital Humanities & NLP'
    """
    t = normalize(text)
    primary = secondary = ""
    keywords: list[str] = []

    has_explicit = re.search(r"Primary Research Area:", t, re.IGNORECASE) is not None

    if has_explicit:
        m = re.search(r"Primary Research Area:\s*(.+?)(?=(?:Secondary Research Area:|" + KEYWORDS_LABEL + r"|$))", t, re.DOTALL | re.IGNORECASE)
        if m:
            primary = re.sub(r"\s+", " ", m.group(1)).strip()
        m = re.search(r"Secondary Research Area:\s*(.+?)(?=(?:" + KEYWORDS_LABEL + r"|$))", t, re.DOTALL | re.IGNORECASE)
        if m:
            secondary = re.sub(r"\s+", " ", m.group(1)).strip()
        m = re.search(KEYWORDS_LABEL + r"\s*(.+)$", t, re.DOTALL)
        if m:
            kw_raw = re.sub(r"\s+", " ", m.group(1)).strip().rstrip(",")
            keywords = [k.strip() for k in re.split(r"[,;]", kw_raw) if k.strip()]
    else:
        # Areas without labels. Try newlines first (one area per line); fall back to / or ,.
        line_parts = [ln.strip() for ln in t.split("\n") if ln.strip()]
        if len(line_parts) >= 2:
            parts = line_parts
        else:
            parts = [p.strip() for p in re.split(r"[/,;]", re.sub(r"\s+", " ", t)) if p.strip()]
        if parts:
            primary = parts[0]
            if len(parts) > 1:
                secondary = parts[1]
            if len(parts) > 2:
                keywords = parts[2:]
    return primary, secondary, keywords


def is_template_slide(name: str, give: list[str]) -> bool:
    if name.lower().startswith("jana sample"):
        return True
    if any(b.lower().startswith("strength/expertise") for b in give):
        return True
    return False


def pick_headshot(pictures: list) -> object | None:
    """Pick the picture whose aspect ratio is closest to 1 (square)."""
    if not pictures:
        return None
    def aspect_distance(p):
        w, h = p.width, p.height
        if w == 0 or h == 0:
            return float("inf")
        a = w / h
        return abs(a - 1.0)
    return min(pictures, key=aspect_distance)


def find_override(slug: str) -> Path | None:
    """Return the override file for a slug, if one exists."""
    if not PHOTO_OVERRIDES.exists():
        return None
    for ext in ("jpg", "jpeg", "png", "webp"):
        candidate = PHOTO_OVERRIDES / f"{slug}.{ext}"
        if candidate.exists():
            return candidate
    return None


def save_photo(picture, slug: str) -> str | None:
    PHOTO_DIR.mkdir(parents=True, exist_ok=True)
    out_path = PHOTO_DIR / f"{slug}.jpg"

    # Override wins over whatever the heuristic picked from the slide.
    override = find_override(slug)
    if override is not None:
        try:
            img = Image.open(override)
            img.thumbnail((512, 512))
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.save(out_path, "JPEG", quality=85)
            return str(out_path.relative_to(ROOT / "public"))
        except Exception:
            pass

    if picture is None:
        return None
    try:
        blob = picture.image.blob
    except Exception:
        return None
    try:
        img = Image.open(BytesIO(blob))
        img.thumbnail((512, 512))
        if img.mode != "RGB":
            img = img.convert("RGB")
        img.save(out_path, "JPEG", quality=85)
    except Exception:
        # Fall back to writing the raw blob with original extension if Pillow fails.
        ext = (picture.image.ext or "png").lstrip(".")
        out_path = PHOTO_DIR / f"{slug}.{ext}"
        out_path.write_bytes(blob)
    return str(out_path.relative_to(ROOT / "public"))


def classify_frame(raw: str) -> str:
    """Return one of: 'give', 'get', 'areas-explicit', 'email', 'name', 'other'."""
    normalized = normalize(raw)
    section, _ = parse_give_get_block(raw)
    if section == "give":
        return "give"
    if section == "get":
        return "get"
    if re.search(r"Primary Research Area:", normalized, re.IGNORECASE):
        return "areas-explicit"
    if re.search(KEYWORDS_LABEL, normalized):
        return "areas-explicit"
    if EMAIL_RE.search(normalized):
        return "email"
    lines = [ln for ln in normalized.split("\n") if ln.strip()]
    # The name frame is always a single line: "Name (location)".
    if len(lines) == 1 and NAME_LOC_RE.match(lines[0].strip()):
        return "name"
    return "other"


def parse_slide(slide, idx: int) -> Researcher | None:
    name = ""
    location = ""
    give: list[str] = []
    get: list[str] = []
    primary = secondary = ""
    keywords: list[str] = []
    email = ""
    links: list[str] = []
    pictures = []

    text_frames: list[str] = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(shape)
        elif shape.has_text_frame:
            t = shape.text_frame.text
            if t.strip():
                text_frames.append(t)

    classified = [(classify_frame(t), t) for t in text_frames]

    for kind, raw in classified:
        normalized = normalize(raw)
        if kind == "give":
            _, bullets = parse_give_get_block(raw)
            give.extend(bullets)
        elif kind == "get":
            _, bullets = parse_give_get_block(raw)
            get.extend(bullets)
        elif kind == "areas-explicit":
            p, s, kw = parse_areas_block(raw)
            if p: primary = p
            if s: secondary = s
            if kw: keywords = kw
        elif kind == "email":
            m = EMAIL_RE.search(normalized)
            if m:
                email = m.group(0)
            links.extend(URL_RE.findall(normalized))
        elif kind == "name":
            first_line = normalized.split("\n", 1)[0].strip()
            m = NAME_LOC_RE.match(first_line)
            if m and not name:
                name = m.group("name").strip()
                location = normalize_location(m.group("loc").strip())

    # Fallback: any unclassified text frame (kind == 'other') likely contains the areas
    # for slides that omit the "Primary Research Area:" labels.
    if not primary:
        for kind, raw in classified:
            if kind == "other":
                p, s, kw = parse_areas_block(raw)
                if p:
                    primary, secondary, keywords = p, s or secondary, kw or keywords
                    break

    if not name:
        return None
    if is_template_slide(name, give):
        return None

    slug = slugify(name)
    head = pick_headshot(pictures)
    # Always go through save_photo so the override mechanism wins even when the
    # slide has no usable image at all.
    photo = save_photo(head, slug)

    return Researcher(
        id=slug,
        name=name,
        location=location,
        give=give,
        get=get,
        primary_area=primary,
        secondary_area=secondary,
        keywords=keywords,
        email=email,
        links=links,
        photo=photo,
        slide_index=idx,
    )


def main() -> int:
    if not PPTX_PATH.exists():
        print(f"PPTX not found at {PPTX_PATH}", file=sys.stderr)
        return 1

    prs = Presentation(str(PPTX_PATH))
    researchers: list[Researcher] = []
    for idx, slide in enumerate(prs.slides):
        if idx == 0:
            continue  # overview slide
        r = parse_slide(slide, idx)
        if r is None:
            continue
        researchers.append(r)

    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps([asdict(r) for r in researchers], indent=2, ensure_ascii=False))
    print(f"Parsed {len(researchers)} researchers → {OUT_JSON}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
